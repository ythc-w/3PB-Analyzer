"""Regression checks for the v4.3 Yield methods and Excel table formatting."""
from pathlib import Path
from tempfile import TemporaryDirectory

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import three_point_bending_analysis as a
from bending_plotting import PlotOptions


def synthetic_curve():
    x=np.linspace(0,1.0,101)
    y=np.zeros_like(x)
    # preload-like low-force region
    y[:10]=np.linspace(0,0.2,10)
    # elastic region, then progressive softening
    for i,xx in enumerate(x[10:],10):
        if xx <= .55:
            y[i]=100*(xx-.08)
        elif xx <= .78:
            y[i]=47 + 42*(xx-.55)
        else:
            y[i]=56.66 - 35*(xx-.78)
    # terminal fracture drop
    y[-3:]=[10,2,0]
    return pd.DataFrame({'Displacement_mm':x,'Force_N':y})


def run():
    raw=synthetic_curve()
    meta=a.SampleMeta('synthetic','synthetic.csv',thickness_mm=1.2,span_mm=7.0)

    # Published 3PB family works without geometry and uses both parameters.
    nogeo=a.replace(meta,thickness_mm=None,span_mm=None)
    r0,d0=a.analyze_one_sample(nogeo,raw,False,yield_method='three_pb',
                               max_displacement_offset_fraction=.002,
                               stiffness_loss_fraction=0.0)
    assert np.isfinite(r0.yield_force_n) and np.isfinite(r0.yield_disp_mm)
    assert d0['yield_method']=='three_pb'
    assert np.isclose(d0['yield_force_constant'],1.0)
    assert np.isclose(d0['delta_offset'][0],d0['recorded_max_displacement_mm']*.002)

    r10,d10=a.analyze_one_sample(nogeo,raw,False,yield_method='three_pb',
                                 max_displacement_offset_fraction=.002,
                                 stiffness_loss_fraction=.10)
    assert np.isclose(d10['yield_force_constant'],.9)
    assert not (np.isclose(r10.yield_force_n,r0.yield_force_n) and np.isclose(r10.yield_disp_mm,r0.yield_disp_mm))

    # Flexural strain-offset requires h and L and keeps the offset independently adjustable.
    rs,ds=a.analyze_one_sample(meta,raw,True,yield_method='strain',offset_strain=.003)
    expected=.003*meta.span_mm**2/(6*meta.thickness_mm)
    assert np.isclose(ds['delta_offset'][0],expected)
    assert rs.fracture_strain_pct is not None
    try:
        a.analyze_one_sample(nogeo,raw,False,yield_method='strain',offset_strain=.002)
    except ValueError:
        pass
    else:
        raise AssertionError('strain-offset Yield accepted missing h/L')

    # Post-yield displacement is a separate interval.
    assert np.isclose(r0.postyield_disp_mm,r0.fracture_disp_mm-r0.yield_disp_mm)

    # Excel tables: visible gridlines, explicit borders, horizontal and vertical center alignment.
    with TemporaryDirectory() as td:
        out=Path(td)
        a.write_outputs([r0],{r0.sample_id:raw},{r0.sample_id:d0},out,False,PlotOptions())
        wb=load_workbook(out/'three_point_bending_results.xlsx')
        for ws in wb.worksheets:
            assert ws.sheet_view.showGridLines is True
            populated=[c for row in ws.iter_rows() for c in row if c.value is not None]
            assert populated
            assert all(c.alignment.horizontal=='center' and c.alignment.vertical=='center' for c in populated)
            assert all(c.border.left.style=='thin' and c.border.right.style=='thin' and
                       c.border.top.style=='thin' and c.border.bottom.style=='thin' for c in populated)

    print('PASS v4.3: 3PB combined parameters, strain offset, PYD, and centered bordered Excel tables')


if __name__=='__main__':
    run()
